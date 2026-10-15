"""Konwersje dla laboratoriów 2 i 3. Bez interfejsu graficznego."""

import csv
import json
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from collections.abc import Callable


class BladKonwersji(Exception):
    """Błąd zrozumiały dla osoby korzystającej z konwertera."""


Konwerter = Callable[[Path, Path], Path]
REJESTR: dict[tuple[str, str], Konwerter] = {}


def konwerter(wejscie: str, wyjscie: str):
    """Rejestruje konwersję, np. ('.csv', '.json')."""
    def zarejestruj(funkcja: Konwerter) -> Konwerter:
        REJESTR[(wejscie, wyjscie)] = funkcja
        return funkcja
    return zarejestruj


def dostepne_formaty(plik: Path) -> list[str]:
    """Zwraca formaty możliwe do uzyskania z danego pliku."""
    return sorted(cel for (zrodlo, cel) in REJESTR if zrodlo == plik.suffix.lower())


def wolna_nazwa(plik: Path, katalog: Path, rozszerzenie: str) -> Path:
    """Nie nadpisuje wcześniejszych wyników; dopisuje _1, _2 itd."""
    wynik = katalog / (plik.stem + rozszerzenie)
    numer = 1
    while wynik.exists():
        wynik = katalog / f"{plik.stem}_{numer}{rozszerzenie}"
        numer += 1
    return wynik


@dataclass(frozen=True)
class Zlecenie:
    plik: Path
    format_docelowy: str
    katalog: Path

    def sprawdz(self) -> None:
        if not self.plik.is_file():
            raise BladKonwersji(f"Nie znaleziono pliku: {self.plik}")
        if not self.katalog.is_dir():
            raise BladKonwersji(f"Nie znaleziono katalogu: {self.katalog}")
        para = (self.plik.suffix.lower(), self.format_docelowy.lower())
        if para not in REJESTR:
            raise BladKonwersji(
                f"Nieobsługiwana konwersja {para[0]} -> {para[1]}. "
                f"Dostępne formaty: {', '.join(dostepne_formaty(self.plik)) or 'brak'}"
            )


def wykonaj(zlecenie: Zlecenie) -> Path:
    zlecenie.sprawdz()
    para = (zlecenie.plik.suffix.lower(), zlecenie.format_docelowy.lower())
    try:
        return REJESTR[para](zlecenie.plik, zlecenie.katalog)
    except (OSError, UnicodeError, PermissionError) as blad:
        raise BladKonwersji(f"Błąd odczytu lub zapisu: {blad}") from blad


# LABORATORIUM 2: CSV <-> JSON
# Schemat JSON: {"columns": ["Imie", "Miasto"],
#                "rows": [{"Imie": "Ala", "Miasto": "Łódź"}]}
# Wszystkie wartości komórek są napisami, dzięki czemu np. "0012" nie staje się 12.

@konwerter(".csv", ".json")
def csv_na_json(plik: Path, katalog: Path) -> Path:
    with plik.open("r", encoding="utf-8-sig", newline="") as wejscie:
        czytnik = csv.reader(wejscie, delimiter=",")
        naglowki = next(czytnik, None)
        if not naglowki or any(not n for n in naglowki):
            raise BladKonwersji("CSV musi zawierać w pierwszym wierszu niepuste nagłówki.")
        if len(set(naglowki)) != len(naglowki):
            raise BladKonwersji("CSV zawiera powtarzające się nazwy kolumn.")

        rekordy = []
        for numer, wiersz in enumerate(czytnik, start=2):
            if len(wiersz) != len(naglowki):
                raise BladKonwersji(
                    f"CSV: wiersz {numer} ma {len(wiersz)} pól, "
                    f"a nagłówek ma {len(naglowki)} kolumn."
                )
            rekordy.append(dict(zip(naglowki, wiersz)))

    wynik = wolna_nazwa(plik, katalog, ".json")
    dane = {"columns": naglowki, "rows": rekordy}
    with wynik.open("w", encoding="utf-8") as wyjscie:
        json.dump(dane, wyjscie, ensure_ascii=False, indent=2)
    return wynik


@konwerter(".json", ".csv")
def json_na_csv(plik: Path, katalog: Path) -> Path:
    try:
        with plik.open("r", encoding="utf-8-sig") as wejscie:
            dane = json.load(wejscie)
    except json.JSONDecodeError as blad:
        raise BladKonwersji(
            f"Niepoprawny JSON (linia {blad.lineno}, kolumna {blad.colno}): {blad.msg}"
        ) from blad

    if not isinstance(dane, dict) or set(dane) != {"columns", "rows"}:
        raise BladKonwersji("JSON musi mieć pola 'columns' (lista nagłówków) i 'rows' (lista rekordów).")

    naglowki, rekordy = dane["columns"], dane["rows"]
    if (not isinstance(naglowki, list) or not naglowki
            or any(not isinstance(n, str) or not n for n in naglowki)
            or len(set(naglowki)) != len(naglowki)):
        raise BladKonwersji("'columns' musi zawierać unikalne, niepuste nazwy kolumn.")
    if not isinstance(rekordy, list):
        raise BladKonwersji("'rows' musi być listą rekordów.")

    for numer, rekord in enumerate(rekordy, start=1):
        if not isinstance(rekord, dict) or set(rekord) != set(naglowki):
            raise BladKonwersji(f"Rekord {numer}: brakujące, nadmiarowe lub nieprawidłowe kolumny.")
        if any(not isinstance(rekord[n], str) for n in naglowki):
            raise BladKonwersji(f"Rekord {numer}: wartości komórek muszą być tekstem.")

    wynik = wolna_nazwa(plik, katalog, ".csv")
    # utf-8-sig dodaje BOM, aby Excel na Windows łatwiej rozpoznawał polskie znaki.
    with wynik.open("w", encoding="utf-8-sig", newline="") as wyjscie:
        pisarz = csv.writer(wyjscie, delimiter=",")
        pisarz.writerow(naglowki)
        for rekord in rekordy:
            pisarz.writerow([rekord[n] for n in naglowki])
    return wynik


# LABORATORIUM 3: TXT / DOCX

@konwerter(".txt", ".docx")
def txt_na_docx(plik: Path, katalog: Path) -> Path:
    try:
        from docx import Document
    except ImportError as blad:
        raise BladKonwersji("Brak biblioteki python-docx. Zainstaluj: pip install python-docx") from blad

    tekst = plik.read_text(encoding="utf-8-sig")
    dokument = Document()
    for linia in tekst.splitlines():
        dokument.add_paragraph(linia)
    wynik = wolna_nazwa(plik, katalog, ".docx")
    dokument.save(wynik)
    return wynik


@konwerter(".docx", ".txt")
def docx_na_txt(plik: Path, katalog: Path) -> Path:
    try:
        from docx import Document
    except ImportError as blad:
        raise BladKonwersji("Brak biblioteki python-docx. Zainstaluj: pip install python-docx") from blad

    try:
        dokument = Document(plik)
        tekst = "\n".join(akapit.text for akapit in dokument.paragraphs)
    except Exception as blad:
        raise BladKonwersji(f"Nie można odczytać DOCX (plik może być uszkodzony): {blad}") from blad

    wynik = wolna_nazwa(plik, katalog, ".txt")
    wynik.write_text(tekst, encoding="utf-8")
    return wynik


# LABORATORIUM 3: DOCX / ODT -> PDF z użyciem LibreOffice

def znajdz_libreoffice() -> str:
    """Znajduje LibreOffice z PATH, zmiennej lub standardowych lokalizacji."""
    wskazany = os.environ.get("LIBREOFFICE_PATH")
    if wskazany:
        znaleziony = shutil.which(wskazany)
        if znaleziony:
            return znaleziony
        raise BladKonwersji(f"Nie znaleziono LibreOffice pod LIBREOFFICE_PATH={wskazany}")

    for nazwa in ("soffice", "libreoffice", "soffice.exe"):
        znaleziony = shutil.which(nazwa)
        if znaleziony:
            return znaleziony

    for kandydat in (
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
    ):
        if kandydat.is_file():
            return str(kandydat)
    raise BladKonwersji(
        "Nie znaleziono LibreOffice. Zainstaluj LibreOffice i dodaj soffice do PATH "
        "lub ustaw zmienną środowiskową LIBREOFFICE_PATH."
    )


def przez_libreoffice(plik: Path, katalog: Path, format_docelowy: str) -> Path:
    program = znajdz_libreoffice()
    wynik = wolna_nazwa(plik, katalog, f".{format_docelowy}")

    # Tymczasowy katalog zapobiega kolizjom nazw z istniejącymi plikami.
    with tempfile.TemporaryDirectory(prefix="konwerter_", dir=katalog) as tymczasowy:
        temp = Path(tymczasowy)
        komenda = [
            program,
            f"-env:UserInstallation={(temp / 'profil').as_uri()}",
            "--headless",
            "--convert-to", format_docelowy,
            "--outdir", str(temp),
            str(plik.resolve()),
        ]
        try:
            proces = subprocess.run(
                komenda,
                shell=False,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=90,
            )
        except subprocess.TimeoutExpired as blad:
            raise BladKonwersji("LibreOffice przekroczył limit 90 sekund.") from blad
        except OSError as blad:
            raise BladKonwersji(f"Nie udało się uruchomić LibreOffice: {blad}") from blad

        if proces.returncode != 0:
            szczegoly = (proces.stderr or proces.stdout).strip()
            raise BladKonwersji(f"LibreOffice zakończył działanie z kodem {proces.returncode}: {szczegoly}")

        utworzony = temp / f"{plik.stem}.{format_docelowy}"
        if not utworzony.is_file():
            szczegoly = (proces.stderr or proces.stdout).strip()
            raise BladKonwersji(f"LibreOffice nie utworzył oczekiwanego pliku. {szczegoly}")
        shutil.move(str(utworzony), str(wynik))
    return wynik


@konwerter(".docx", ".pdf")
def docx_na_pdf(plik: Path, katalog: Path) -> Path:
    return przez_libreoffice(plik, katalog, "pdf")


@konwerter(".odt", ".pdf")
def odt_na_pdf(plik: Path, katalog: Path) -> Path:
    return przez_libreoffice(plik, katalog, "pdf")
